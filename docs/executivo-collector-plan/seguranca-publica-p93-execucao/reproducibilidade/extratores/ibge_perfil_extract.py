"""MUNIC e ESTADIC, edições 2019 e 2023 (IBGE, Pesquisas de Informações Básicas): aba "Segurança Pública" + dicionário das variáveis usadas.
Uso: python -I scripts/exec/ibge_perfil_extract.py munic|estadic <xlsx> <saida.json> <edicao>

Lê somente as colunas necessárias (códigos e variáveis de segurança) e o rótulo oficial de cada variável no dicionário, para que o publicador
confira o layout (um rótulo diferente do esperado aborta a carga). Sem dados pessoais: unidades territoriais e agregados institucionais.
"""
import json, re, sys, unicodedata
import openpyxl

VERSION = 'ibge-perfil-seg-v2'
MUNIC_VARS = ['MSEG27', 'MSEG281', 'MSEG282', 'MSEG283'] + [f'MSEG16{i}' for i in range(1, 10)] + ['MSEG1610', 'MSEG1611', 'MSEG17', 'MSEG18', 'MSEG19']
ESTADIC_VARS = ['ESEG271', 'ESEG272', 'ESEG273', 'ESEG491', 'ESEG492', 'ESEG493', 'ESEG3011', 'ESEG3021', 'ESEG3031', 'ESEG5111', 'ESEG5114', 'ESEG5121', 'ESEG5124', 'ESEG5131', 'ESEG5134', 'ESEG50', 'ESEG082', 'ESEG081']
# core = variáveis que a edição 2019 também tem (mesmo código e mesmo rótulo); as demais são opcionais e só existem na edição 2023.
MUNIC_CORE = ['MSEG27', 'MSEG281', 'MSEG282', 'MSEG283', 'MSEG161', 'MSEG17', 'MSEG18', 'MSEG19']
MODES = {
    'munic': {'vars': MUNIC_VARS, 'core': MUNIC_CORE, 'code': ['CODMUN'], 'uf': ['CODUF'], 'expect_rows': 5570},
    'estadic': {'vars': ESTADIC_VARS, 'core': ESTADIC_VARS, 'code': [], 'uf': ['CODUF'], 'expect_rows': 27},
}

def norm(s): return ''.join(c for c in unicodedata.normalize('NFKD', str(s)) if not unicodedata.combining(c)).upper().strip()
def key(s): return re.sub(r'[^A-Z0-9]', '', norm(s))

def sheet_named(wb, *needles):
    for ws in wb.worksheets:
        n = norm(ws.title)
        if all(x in n for x in needles): return ws
    raise SystemExit(f'aba não encontrada: {needles}; abas: {[w.title for w in wb.worksheets]}')

def dictionary(wb, wanted):
    ws = sheet_named(wb, 'DICION'); out = {}; head = ''; sub = ''
    for r in ws.iter_rows(values_only=True):
        cells = [('' if c is None else str(c).replace('\n', ' ').strip()) for c in r[:8]]
        code = next((c for c in cells if c in wanted), None)
        if code:
            label = ' / '.join(c for c in cells if c and c != code)
            out[code] = {'label': label, 'context': ' | '.join(x for x in (head, sub) if x)}
        else:
            txt = [c for c in cells if c]
            if txt and cells[0]: head, sub = cells[0], ''
            elif txt and cells[1] and not cells[2]: sub = cells[1]
    return out

def main(mode, xlsx, out, edition):
    cfg = MODES[mode]
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = sheet_named(wb, 'SEGURANCA', 'PUBLICA')
    rows = ws.iter_rows(values_only=True); head = [str(c) if c is not None else '' for c in next(rows)]
    ix = {h: i for i, h in enumerate(head)}; kx = {key(h): i for i, h in enumerate(head)}
    def col(names):
        for n in names:
            if n in kx: return kx[n]
        raise SystemExit(f'coluna de identificação ausente ({names}) na aba {ws.title}')
    ic = col(cfg['code']) if cfg['code'] else None; iu = col(cfg['uf'])
    missing_core = [v for v in cfg['core'] if v not in ix]
    if missing_core: raise SystemExit(f'variáveis obrigatórias ausentes na aba {ws.title}: {missing_core}')
    present = [v for v in cfg['vars'] if v in ix]; missing_optional = [v for v in cfg['vars'] if v not in ix]
    dic = dictionary(wb, set(present))
    if len(dic) != len(present): raise SystemExit(f'dicionário sem as variáveis: {sorted(set(present) - set(dic))}')
    data = []
    for r in rows:
        if r[iu] is None: continue
        rec = {'ufCode': int(r[iu]), 'v': {v: r[ix[v]] for v in present}}
        if ic is not None: rec['cod'] = int(r[ic])
        data.append(rec)
    if len(data) != cfg['expect_rows']: raise SystemExit(f'{len(data)} linhas (esperado {cfg["expect_rows"]})')
    json.dump({'version': VERSION, 'mode': mode, 'edition': int(edition), 'sheet': ws.title, 'columns': len(head), 'rows': data, 'dictionary': dic, 'missingOptional': missing_optional}, open(out, 'w', encoding='utf-8'), ensure_ascii=False)
    print(json.dumps({'mode': mode, 'edition': int(edition), 'rows': len(data), 'vars': len(present), 'missingOptional': missing_optional, 'columns': len(head)}))

if __name__ == '__main__':
    if len(sys.argv) != 5 or sys.argv[1] not in MODES: raise SystemExit(__doc__)
    main(*sys.argv[1:])
