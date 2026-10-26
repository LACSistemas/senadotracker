import hashlib
import json
import re
import sys
from decimal import Decimal, InvalidOperation
import openpyxl

path, out_path = sys.argv[1:3]
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb.active
header = None
for row in ws.iter_rows(values_only=True):
    vals = [str(x).strip() if x is not None else '' for x in row]
    if 'Cod. Escola' in vals:
        header = vals
        break
if header is None:
    raise RuntimeError('PDDE header not found')
idx = {v: i for i, v in enumerate(header) if v}
def val(row, key):
    i = idx.get(key)
    return '' if i is None or i >= len(row) or row[i] is None else str(row[i]).strip()
def money(v):
    if not v: return None
    try:
        return int((Decimal(str(v).replace('.', '').replace(',', '.')) * 100).quantize(Decimal('1')))
    except (InvalidOperation, ValueError):
        return None
def digits(v):
    x = re.sub(r'\D', '', v or '')
    return x or None
rows = []
school = {}
for raw in ws.iter_rows(values_only=True):
    dest = val(raw, 'Destinação')
    if not dest or dest.lower().startswith('subtotal') or dest.lower().startswith('total geral'):
        continue
    for k in ['UF','Município','Rede de Ensino','Cod. Escola','Nome Escola','CNPJ UEx','Unidade Executora','CNPJ EEx/EM','Presta��o de Contas','Dados Cadastrais - PDDEWeb']:
        x = val(raw, k)
        if x: school[k] = x
    if not school.get('Cod. Escola'): continue
    base = {
        'uf': school.get('UF',''), 'municipality_name': school.get('Município'), 'school_inep': digits(school.get('Cod. Escola')),
        'school_name': school.get('Nome Escola'), 'uex_cnpj': digits(school.get('CNPJ UEx')), 'uex_name': school.get('Unidade Executora'),
        'eex_cnpj': digits(school.get('CNPJ EEx/EM')), 'destination': dest, 'status': school.get('Presta��o de Contas'),
        'cadastro_status': school.get('Dados Cadastrais - PDDEWeb'),
    }
    for key, col in [('DUE', 'Vl Final Devido Total'), ('PAID', 'Valor Pago Total')]:
        amount = money(val(raw, col))
        if amount is None: continue
        item = dict(base, financial_stage=key, amount_cents=amount)
        canonical = json.dumps(item, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
        item['raw_id'] = hashlib.sha256(canonical.encode()).hexdigest()
        rows.append(item)
rows.sort(key=lambda x: (x['raw_id'], x['financial_stage']))
with open(out_path, 'w', encoding='utf-8') as out:
    for row in rows: out.write(json.dumps(row, ensure_ascii=False) + '\n')
print(json.dumps({'rows': len(rows), 'output': out_path}))
