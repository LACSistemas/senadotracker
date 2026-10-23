"""Agrega as planilhas abertas da SSPDS-CE em município × mês × natureza (P9.6). SEM atributo individual na saída.

  python -I scripts/exec/ce_sspds_extract.py <cvli|cvp> <arquivo.xlsx> <saida.json>

  cvli  aba "CVLI": uma linha por vítima (Município, AIS, Natureza, Data, ...) — gênero, idade, escolaridade, raça, meio e hora são descartados
  cvp   aba única: uma linha por ocorrência de CVP (AIS, Município, Data, ...); natureza = TOTAL
Saída: {"kind", "header", "rows": [[AAAAMM, município, natureza, contagem]], "lines", "badDates"}.
"""
import json, re, sys, unicodedata
from collections import defaultdict
import openpyxl


def norm(s):
    s = unicodedata.normalize('NFD', str(s if s is not None else ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Za-z0-9 ]+', ' ', s)).strip().upper()


def main():
    kind, src, out = sys.argv[1:4]
    if kind not in ('cvli', 'cvp'):
        raise SystemExit('kind inválido')
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb['CVLI'] if kind == 'cvli' and 'CVLI' in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    head = [norm(c) for c in next(it)]
    def col(name):
        if name not in head:
            raise SystemExit(f'coluna ausente: {name} em {head}')
        return head.index(name)
    mun, dt = col('MUNICIPIO'), col('DATA')
    nat = col('NATUREZA') if kind == 'cvli' else None
    agg = defaultdict(int)
    n = bad = 0
    for r in it:
        if r is None or r[mun] is None:
            continue
        d = r[dt]
        if not hasattr(d, 'year'):
            bad += 1
            continue
        agg[(d.year * 100 + d.month, str(r[mun]).strip(), norm(r[nat]) if nat is not None else 'TOTAL')] += 1
        n += 1
    rows = [[k[0], k[1], k[2], v] for k, v in sorted(agg.items())]
    json.dump({'kind': kind, 'header': head, 'rows': rows, 'lines': n, 'badDates': bad}, open(out, 'w', encoding='utf-8'), ensure_ascii=False)


if __name__ == '__main__':
    main()
