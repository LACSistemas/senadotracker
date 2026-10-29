import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

SOURCE = Path(r"D:\Civica\raw\bem-estar\cotidiano\municip-2024\Base_MUNIC_2024_20251107.xlsx")
OUT = Path(r"D:\Civica\staging\bem-estar\cotidiano\p126_municip_rows.json")

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()

def clean(v):
    if v is None:
        return None
    return str(v).strip()

def main():
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = next(s for s in wb.worksheets if s.title.startswith("Inform"))
    it = ws.iter_rows(values_only=True)
    headers = [clean(v) for v in next(it)]
    idx = {h.lower(): i for i, h in enumerate(headers) if h}
    required = ["Cod Munic", "Uf", "Mtic092", "Mtic10", "Mtic165", "Mtic22", "Mtic241"]
    missing = [x for x in required if x.lower() not in idx]
    if missing:
        raise RuntimeError(f"missing columns: {missing}")

    rows = []
    stats = {"sourceRows": 0, "mappedRows": 0, "unresolvedRows": 0, "byUf": {}, "indicators": {}}
    for row in it:
        stats["sourceRows"] += 1
        code = clean(row[idx["cod munic" ]])
        if not code or not code.isdigit():
            stats["unresolvedRows"] += 1
            continue
        ibge7 = int(code)
        uf = clean(row[idx["uf"]])
        stats["mappedRows"] += 1
        stats["byUf"][uf] = stats["byUf"].get(uf, 0) + 1
        values = {
            "digital_service": clean(row[idx["mtic092"]]),
            "website_active": clean(row[idx["mtic10"]]),
            "mobile_website": clean(row[idx["mtic165"]]),
            "wifi_public": clean(row[idx["mtic22"]]),
            "wifi_whole_municipality": clean(row[idx["mtic241"]]),
        }
        rows.append({"ibge7": ibge7, "uf": uf, **values})
        for key, val in values.items():
            stats["indicators"].setdefault(key, {}).setdefault(val, 0)
            stats["indicators"][key][val] += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "sourceFiles": [{
            "provider": "IBGE",
            "family": "MUNIC",
            "edition": "2024",
            "resource": "Base_MUNIC_2024_20251107.xlsx",
            "url": "https://ftp.ibge.gov.br/Perfil_Municipios/2024/Base_de_Dados/Base_MUNIC_2024_20251107.xlsx",
            "path": str(SOURCE),
            "bytes": SOURCE.stat().st_size,
            "sha256": sha256(SOURCE),
            "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "sheet": ws.title,
            "httpStatus": 200,
        }],
        "edition": 2024,
        "sheet": ws.title,
        "rows": rows,
        "stats": stats,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"out": str(OUT), "bytes": OUT.stat().st_size, "stats": stats}, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
