"""P10.4 — extrai as planilhas oficiais de resultados do Saeb (agregados por Brasil, UF e município; nenhum microdado) para um TSV LONGO:
uma linha por (território × dependência × localização × capital × ano/série × disciplina) com média de proficiência e/ou distribuição por nível.
Uso: python -I -X utf8 saeb_extract.py <ano> <nivel> <saida.tsv> <arquivo> [<aba>]      nivel = brasil | regioes | estados | municipios
 - Arquivos: .xlsx (2019, 2021, 2025), .xlsb (2023, extraído do .rar) e .xlsx de dentro do ZIP de 2017 (TS_BRASIL, TS_REGIAO, TS_UF, TS_MUNICIPIO). Cada aba/arquivo traz as colunas
   MEDIA_<série>_<LP|MT|CH|CN> e nivel_<n>_<LP|MT|CH|CN><série>; a série é 2, 5, 9 (ensino fundamental) e 12, 13, 14 (ensino médio, conforme a edição). 2021 tem uma linha de descrição logo abaixo do cabeçalho (ignorada).
 - Alfabetização: TX_ALFABETIZADO (2021) e PC_ALFABETIZADO (2023) vão para a linha 2º ano / LP da mesma chave; não se mistura com a nota.
 - "-", vazio e texto são ausentes (nunca zero). Nada é recalculado ou ponderado.
Saída: row year level uf ibge7 name dep loc cap grade subj media tx_alf nv0..nv10. Última linha de stdout: JSON de resumo."""
import json, os, re, sys

MEDRX = re.compile(r'^MEDIA_(\d+)_(LP|MT|CH|CN)$', re.I)
NIVRX = re.compile(r'^nivel_(\d+)_(LP|MT|CH|CN)(\d+)$', re.I)
ALF = ('TX_ALFABETIZADO', 'PC_ALFABETIZADO')
NIV = [f'nv{i}' for i in range(11)]


def sheet_rows(path, sheet):
    if path.lower().endswith('.xlsb'):
        import pyxlsb
        with pyxlsb.open_workbook(path) as wb:
            names = wb.sheets
            with wb.get_sheet(sheet if sheet in names else names[0]) as sh:
                for r in sh.rows(): yield [c.v for c in r]
        return
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    for r in ws.iter_rows(values_only=True): yield list(r)


def num(v):
    if v is None: return ''
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        v = round(float(v), 6)   # ruído de ponto flutuante do Excel (15.280000000000001)
        return repr(v) if v != int(v) else str(int(v))
    s = str(v).strip().replace(',', '.')
    return s if re.fullmatch(r'-?\d+(\.\d+)?', s) else ''


def main():
    year, level, out, path = int(sys.argv[1]), sys.argv[2], sys.argv[3], sys.argv[4]
    sheet = sys.argv[5] if len(sys.argv) > 5 else {'brasil': 'Brasil', 'regioes': 'Regiões', 'estados': 'Estados', 'municipios': 'Municípios'}[level]
    it = sheet_rows(path, sheet); head = [str(c).strip() if c is not None else '' for c in next(it)]
    ix = {c: i for i, c in enumerate(head)}
    meds = {(int(m.group(1)), m.group(2).upper()): i for i, c in enumerate(head) if (m := MEDRX.match(c))}
    nivs = {}
    for i, c in enumerate(head):
        m = NIVRX.match(c)
        if m: nivs.setdefault((int(m.group(3)), m.group(2).upper()), {})[int(m.group(1))] = i
    cols = set(meds) | set(nivs)
    alf = next((ix[a] for a in ALF if a in ix), None)
    def g(r, name): return str(r[ix[name]]).strip() if name in ix and ix[name] < len(r) and r[ix[name]] is not None else ''
    n = 0; skipped = 0; maxlev = 0
    with open(out + '.part', 'w', encoding='utf-8', newline='') as o:
        o.write('row\tyear\tlevel\tuf\tibge7\tname\tdep\tloc\tcap\tgrade\tsubj\tmedia\ttx_alf\t' + '\t'.join(NIV) + '\n')
        for r in it:
            if not r or all(c is None for c in r): continue
            ano = g(r, 'ANO_SAEB')
            if 'ANO_SAEB' in ix and not re.fullmatch(r'\d{4}(\.0)?', ano):
                skipped += 1; continue                                   # linha de descrição (2021) ou lixo
            if 'ANO_SAEB' in ix and int(float(ano)) != year: continue
            uf = g(r, 'CO_UF'); ib = g(r, 'CO_MUNICIPIO'); nm = g(r, 'NO_MUNICIPIO') or g(r, 'NO_UF') or g(r, 'NO_REGIAO') or g(r, 'ID')
            uf = re.sub(r'\.0$', '', uf); ib = re.sub(r'\.0$', '', ib)
            dep, loc, cap = g(r, 'DEPENDENCIA_ADM'), g(r, 'LOCALIZACAO'), g(r, 'CAPITAL')
            for (gr, su) in sorted(cols):
                med = num(r[meds[(gr, su)]]) if (gr, su) in meds and meds[(gr, su)] < len(r) else ''
                lv = nivs.get((gr, su), {}); nv = [''] * 11
                for k, i in lv.items():
                    if k <= 10 and i < len(r): nv[k] = num(r[i])
                    maxlev = max(maxlev, k)
                ta = num(r[alf]) if alf is not None and gr == 2 and su == 'LP' and alf < len(r) else ''
                if med == '' and ta == '' and not any(nv): continue
                n += 1
                o.write('\t'.join([str(n), str(year), level, uf, ib, nm.replace('\t', ' '), dep, loc, cap, str(gr), su, med, ta] + nv) + '\n')
    os.replace(out + '.part', out)
    print(json.dumps({'year': year, 'level': level, 'file': os.path.basename(path), 'sheet': sheet, 'rows': n, 'skipped_description_rows': skipped, 'series_subjects': sorted(f'{a}{b}' for a, b in cols), 'alfabetizacao_col': head[alf] if alf is not None else None, 'max_level_index': maxlev}))


main()
