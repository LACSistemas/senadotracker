from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(r"D:\Civica\raw\bem-estar\censo-suas\2024\github")
OUT = Path(r"D:\Civica\staging\bem-estar\censo-suas\2024\p124r_rows.json")

def pick(pattern: str) -> Path:
    xs = list(BASE.glob(pattern))
    if len(xs) != 1:
        raise RuntimeError(f"pattern {pattern!r}: {len(xs)} files")
    return xs[0]

def sha(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest().upper()

def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None

def code(v, width: int) -> str | None:
    s = clean(v)
    if s is None:
        return None
    s = re.sub(r"\.0+$", "", s)
    digits = re.sub(r"\D", "", s)
    return digits.zfill(width) if digits else None

def number(v):
    s = clean(v)
    if s is None or s in {"-", "NA", "nan"}:
        return None
    s = s.replace(".", "").replace(",", ".") if isinstance(v, str) and "," in s else s
    try:
        return float(s)
    except ValueError:
        return None

def sheet_rows(path: Path, sheet: str, header_row: int):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(rows)
    headers = [clean(x) or f"col{i+1}" for i, x in enumerate(next(rows))]
    out = []
    for row in rows:
        if not any(x is not None and str(x).strip() != "" for x in row):
            continue
        out.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    wb.close()
    return out, headers

def sheet_values(path: Path, sheet: str, first_data_row: int):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    out = []
    for row in ws.iter_rows(min_row=first_data_row, values_only=True):
        if any(x is not None and str(x).strip() != "" for x in row):
            out.append(list(row))
    wb.close()
    return out

def main():
    csv_path = pick("Censo_SUAS_2024_CRAS_Dados_Gerais.csv")
    dict_path = pick("Censo_SUAS_2024_CRAS_Dicionario*")
    readme_path = pick("Censo_SUAS_2024_CRAS_Leia*")
    script_path = BASE / "_carregamento.R"
    idcras_path = pick("IDCRAS_*.xlsx")
    idcreas_path = pick("IDCREAS*.xlsx")
    idcon_path = pick("IDConselho*.xlsx")
    files = [csv_path, dict_path, readme_path, script_path, idcras_path, idcreas_path, idcon_path]

    cras = []
    with csv_path.open("r", encoding="latin1", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            ident = clean(r.get("NU_IDENTIFICADOR"))
            ibge7 = code(r.get("IBGE7"), 7)
            if not ident or not ibge7:
                raise RuntimeError("CRAS row without NU_IDENTIFICADOR/IBGE7")
            cras.append({
                "id": ident,
                "ibge7": ibge7,
                "ibge6": ibge7[:6],
                "uf": clean(r.get("UF")),
                "municipio": clean(r.get("Município")),
                "q1": clean(r.get("q1")),
                "q_incompleto": clean(r.get("Q_incompleto")),
            })
    ids = [x["id"] for x in cras]
    if len(cras) != len(set(ids)):
        raise RuntimeError("duplicate CRAS ids")

    idcras_rows, _ = sheet_rows(idcras_path, "IDCRAS_2024_RESULTADO", 7)
    idcras_units = []
    for r in idcras_rows:
        ident = clean(r.get("Nº IDENTIFICADOR do CRAS"))
        ibge7 = code(r.get("IBGE7"), 7)
        if not ident or not ibge7:
            continue
        idcras_units.append({"id": ident, "ibge7": ibge7, "ibge6": ibge7[:6], "score": number(r.get("IDCRAS"))})
    idcras_municipal, _ = sheet_rows(idcras_path, "IDCRAS 2024 MEDIO MUNICIPAL", 7)
    idcras_muni = []
    for r in idcras_municipal:
        ibge7 = code(r.get("IBGE7"), 7)
        if not ibge7:
            continue
        idcras_muni.append({
            "ibge7": ibge7,
            "respondents": number(r.get("Número de CRAS que respondeu ao Censo SUAS 2024")),
            "mean": number(r.get("IDCRAS médio do município")),
        })

    # Use source column positions for IDCREAS: the workbook contains a malformed
    # display header (N� IDENTIFICADOR) in some readers, while the data columns
    # are stable: id, IBGE6, IBGE7, ..., score at column 11.
    idcreas_rows = sheet_values(idcreas_path, "IDCREAS2024_divulgacao", 8)
    idcreas_units = []
    for r in idcreas_rows:
        ident = clean(r[0] if len(r) > 0 else None)
        ibge7 = code(r[2] if len(r) > 2 else None, 7)
        if not ident or not ibge7:
            continue
        idcreas_units.append({"id": ident, "ibge7": ibge7, "ibge6": ibge7[:6], "score": number(r[10] if len(r) > 10 else None)})
    idcreas_municipal = sheet_values(idcreas_path, "IDCREAS 2024 MEDIO MUNICIPAL", 8)
    idcreas_muni = []
    for r in idcreas_municipal:
        ibge6 = code(r[0] if len(r) > 0 else None, 6)
        if not ibge6:
            continue
        idcreas_muni.append({
            "ibge6": ibge6,
            "respondents": number(r[5] if len(r) > 5 else None),
            "mean": number(r[9] if len(r) > 9 else None),
        })

    idcon_rows, _ = sheet_rows(idcon_path, "IDConselho_Municipal_2024_divul", 5)
    idcon = []
    for r in idcon_rows:
        ibge7 = code(r.get("IBGE7"), 7)
        if not ibge7:
            continue
        idcon.append({"ibge7": ibge7, "score": number(r.get("IDConselho"))})

    source_files = []
    for p in files:
        source_files.append({"name": p.name, "relPath": str(p.relative_to(BASE)), "bytes": p.stat().st_size, "sha256": sha(p)})
    out = {
        "generatedAt": __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat(),
        "sourceFiles": source_files,
        "sourceHash": hashlib.sha256(json.dumps(source_files, sort_keys=True).encode()).hexdigest(),
        "cras": {"rows": cras, "units": len(cras), "distinctMunicipalities": len({x["ibge6"] for x in cras})},
        "idcras": {"units": idcras_units, "municipal": idcras_muni},
        "idcreas": {"units": idcreas_units, "municipal": idcreas_muni},
        "idconselho": idcon,
        "counts": {"crasRows": len(cras), "crasIds": len(set(ids)), "idcrasUnits": len(idcras_units), "idcrasMunicipal": len(idcras_muni), "idcreasUnits": len(idcreas_units), "idcreasMunicipal": len(idcreas_muni), "idconselhoMunicipal": len(idcon)},
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(out["counts"], indent=2, ensure_ascii=False))

if __name__ == "__main__":
    main()
