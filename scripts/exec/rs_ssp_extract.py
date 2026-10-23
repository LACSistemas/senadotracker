"""Extrai os indicadores criminais por município da SSP-RS (planilha anual com uma aba por mês) — P9.6.

  python -I scripts/exec/rs_ssp_extract.py <planilha.xlsx> <saida.json>

Só números agregados por município × mês (sem qualquer dado individual). Para cada aba JAN..DEZ: cabeçalho normalizado, linhas por município e a linha "Total RS" (controle da própria fonte).
Também lê, da aba GERAL, o período declarado ("... de 01 de janeiro a 30 de setembro de 2026") e a tabela mensal do estado (controle independente).
"""
import json, re, sys, unicodedata
import openpyxl

MESES = ['JAN', 'FEV', 'MAR', 'ABR', 'MAI', 'JUN', 'JUL', 'AGO', 'SET', 'OUT', 'NOV', 'DEZ']


def norm(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Za-z0-9 ]+', ' ', s)).strip().upper()


def num(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace('.', '').replace(',', '.'))
    except ValueError:
        return None


def main():
    src, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    res = {'sheets': {}, 'geral': None, 'periodo': None}
    names = {n.upper(): n for n in wb.sheetnames}
    if 'GERAL' in names:
        g = list(wb[names['GERAL']].iter_rows(values_only=True))
        for r in g[:8]:
            for c in r:
                if isinstance(c, str) and re.search(r'per.odo de', c, re.I):
                    res['periodo'] = c
        hdr = None
        for i, r in enumerate(g):
            if r and len(r) > 1 and isinstance(r[1], str) and norm(r[1]).startswith('MES OCORRENCIAS'):
                hdr = i
                break
        if hdr is not None:
            head = [norm(c) for c in g[hdr][2:]]
            months = {}
            for r in g[hdr + 1:]:
                lab = r[1] if len(r) > 1 else None
                if isinstance(lab, str) and re.match(r'^\d{4}/', lab):
                    months[lab] = [num(c) for c in r[2:2 + len(head)]]
            res['geral'] = {'header': head, 'months': months}
    for k, m in enumerate(MESES, 1):
        if m not in names:
            continue
        rows = list(wb[names[m]].iter_rows(values_only=True))
        hi = None
        for i, r in enumerate(rows[:15]):
            if r and isinstance(r[0], str) and norm(r[0]).startswith('MUNICIPIOS'):
                hi = i
                break
        if hi is None:
            continue
        head = [norm(c) for c in rows[hi][1:] if c is not None]
        data, total = [], None
        for r in rows[hi + 1:]:
            if not r or r[0] is None or str(r[0]).strip() == '':
                continue
            nm = str(r[0]).strip()
            vals = [num(c) for c in r[1:1 + len(head)]]
            if norm(nm).startswith('TOTAL'):
                total = vals
                break
            data.append([nm] + vals)
        res['sheets'][str(k)] = {'sheet': m, 'header': head, 'rows': data, 'total': total}
    json.dump(res, open(out, 'w', encoding='utf-8'), ensure_ascii=False)


if __name__ == '__main__':
    main()
