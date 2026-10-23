"""Agrega os microdados abertos da SDS-PE em município × mês × natureza (P9.6). SEM atributo individual na saída.

  python -I scripts/exec/pe_sds_extract.py <cvli|cvp|estupro|vd> <arquivo.xlsx> <saida.json>

  cvli    MICRODADOS_DE_CVLI: uma linha por vítima/conjunto (MUNICIPIO, NATUREZA JURIDICA, DATA, TOTAL DE VITIMAS) — sexo, idade e região são descartados
  cvp     Microdados de CVP (crimes violentos contra o patrimônio): já agregado (DATA, MUNICIPIO, TOTAL)
  estupro MICRODADOS_ESTUPRO (MUNICIPIO DO FATO, NATUREZA, DATA DO FATO, TOTAL DE VITIMAS) — sexo e faixa etária descartados
  vd      MICRODADOS_DE_VIOLENCIA_DOMESTICA (layout validado em tempo de execução)
Saída: {"kind", "header", "rows": [[AAAAMM, município, natureza, soma]], "lines", "badDates"}.
"""
import json, re, sys, unicodedata
from collections import defaultdict
import openpyxl


def norm(s):
    s = unicodedata.normalize('NFD', str(s if s is not None else ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Za-z0-9 ]+', ' ', s)).strip().upper()


def pick(h, *names, required=True):
    for n in names:
        for i, c in enumerate(h):
            if c == n:
                return i
    for n in names:
        for i, c in enumerate(h):
            if c.startswith(n):
                return i
    if required:
        raise SystemExit(f'coluna ausente: {names} em {h}')
    return None


def main():
    kind, src, out = sys.argv[1:4]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = max((wb[n] for n in wb.sheetnames), key=lambda w: w.max_row or 0)
    it = ws.iter_rows(values_only=True)
    head = [norm(c) for c in next(it)]
    mun = pick(head, 'MUNICIPIO DO FATO', 'MUNICIPIO')
    dt = pick(head, 'DATA DO FATO', 'DATA')
    if kind == 'cvli':
        nat, val = pick(head, 'NATUREZA JURIDICA', 'NATUREZA'), pick(head, 'TOTAL DE VITIMAS')
    elif kind == 'cvp':
        nat, val = None, pick(head, 'TOTAL')
    elif kind in ('estupro', 'vd'):
        nat, val = pick(head, 'NATUREZA', 'TIPO', required=False), pick(head, 'TOTAL DE VITIMAS', 'TOTAL')
    else:
        raise SystemExit('kind inválido')
    agg = defaultdict(float)
    n = bad = 0
    for r in it:
        if r is None or len(r) <= max(mun, dt, val) or r[mun] is None:
            continue
        d = r[dt]
        if hasattr(d, 'year'):
            ym = d.year * 100 + d.month
        else:
            m = re.match(r'^(\d{4})-(\d{2})', str(d or ''))
            if not m:
                bad += 1
                continue
            ym = int(m.group(1)) * 100 + int(m.group(2))
        v = r[val]
        if v is None:
            continue
        agg[(ym, str(r[mun]).strip(), norm(r[nat]) if nat is not None else 'TOTAL')] += float(v)
        n += 1
    rows = [[k[0], k[1], k[2], v] for k, v in sorted(agg.items())]
    json.dump({'kind': kind, 'header': head, 'rows': rows, 'lines': n, 'badDates': bad}, open(out, 'w', encoding='utf-8'), ensure_ascii=False)


if __name__ == '__main__':
    main()
