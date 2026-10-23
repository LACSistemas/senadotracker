"""Agrega os microdados abertos da SSP-SP (SPDadosCriminais_AAAA.xlsx, um B.O. por linha) em município × mês × natureza (P9.6).

  python -I scripts/exec/sp_ssp_extract.py <arquivo.xlsx> <saida.json>

Lê em streaming (read_only) as abas de dados (cabeçalho com NATUREZA_APURADA, MES_ESTATISTICA, ANO_ESTATISTICA e COD IBGE).
Mantém SÓ ano/mês estatístico, código IBGE, nome do município e natureza apurada; endereço, coordenadas, número de B.O., local e hora são lidos e descartados na própria leitura.
Saída: {"sheets": {aba: linhas}, "rows": [[AAAAMM, ibge7, município, natureza, contagem]], "lines": n, "skipped": {motivo: n}}.
"""
import json, sys, time
from collections import defaultdict
import openpyxl

# nome do campo no contrato → grafias que a SSP-SP usou ao longo dos anos
NEED = {'NOME_MUNICIPIO': ['NOME_MUNICIPIO', 'CIDADE'], 'NATUREZA_APURADA': ['NATUREZA_APURADA'], 'MES_ESTATISTICA': ['MES_ESTATISTICA'], 'ANO_ESTATISTICA': ['ANO_ESTATISTICA'], 'COD IBGE': ['COD IBGE', 'CD_IBGE']}


def locate(h):
    out = {}
    for k, names in NEED.items():
        hit = [h.index(n) for n in names if n in h]
        if not hit:
            raise SystemExit(f'coluna ausente: {names} em {h}')
        out[k] = hit[0]
    return out


def main():
    src, out = sys.argv[1:3]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    agg = defaultdict(int)
    names = {}
    sheets = {}
    skipped = defaultdict(int)
    t0 = time.time()
    for name in wb.sheetnames:
        ws = wb[name]
        it = ws.iter_rows(values_only=True)
        head = next(it, None)
        if not head or 'NATUREZA_APURADA' not in [str(c) for c in head if c]:
            continue
        h = [str(c).strip() if c else '' for c in head]
        ix = locate(h)
        n = 0
        for r in it:
            if r is None or len(r) <= max(ix.values()):
                skipped['curta'] += 1
                continue
            nat, mes, ano, cod = r[ix['NATUREZA_APURADA']], r[ix['MES_ESTATISTICA']], r[ix['ANO_ESTATISTICA']], r[ix['COD IBGE']]
            if nat is None or mes is None or ano is None:
                skipped['sem_natureza_ou_data'] += 1
                continue
            try:
                ym = int(ano) * 100 + int(mes)
                ibge = int(cod)
            except (TypeError, ValueError):
                skipped['sem_ibge'] += 1
                continue
            key = (ym, ibge, str(nat).strip())
            agg[key] += 1
            names.setdefault(ibge, str(r[ix['NOME_MUNICIPIO']] or '').strip())
            n += 1
        sheets[name] = n
    rows = [[k[0], k[1], names.get(k[1], ''), k[2], v] for k, v in sorted(agg.items())]
    json.dump({'sheets': sheets, 'rows': rows, 'lines': sum(sheets.values()), 'skipped': dict(skipped), 'seconds': round(time.time() - t0)}, open(out, 'w', encoding='utf-8'), ensure_ascii=False)


if __name__ == '__main__':
    main()
